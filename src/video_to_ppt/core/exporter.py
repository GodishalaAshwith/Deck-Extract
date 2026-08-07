"""Slide exporter for saving PNG images, JSON metadata, PDF, PPTX, and ZIP archives."""

import json
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional

import cv2
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from .config import ExportConfig
from .detector import DetectedSlide
from .extractor import VideoProbeInfo, format_timestamp_filename


@dataclass
class ExportResult:
    """Paths and summary of exported artifacts."""
    output_dir: Path
    total_slides: int
    image_paths: List[Path]
    metadata_path: Optional[Path] = None
    pdf_path: Optional[Path] = None
    pptx_path: Optional[Path] = None
    zip_path: Optional[Path] = None


class SlideExporter:
    """Exports detected slides to images, PDF, PPTX, and ZIP archives."""

    def __init__(self, config: Optional[ExportConfig] = None):
        self.config = config or ExportConfig()

    def export(
        self,
        slides: List[DetectedSlide],
        probe_info: Optional[VideoProbeInfo] = None
    ) -> ExportResult:
        """Export all detected slides according to configured formats."""
        out_dir = Path(self.config.output_dir).resolve()
        out_dir.mkdir(parents=True, exist_ok=True)

        image_paths: List[Path] = []
        metadata_records = []

        # 1. Save PNG images
        for slide in slides:
            ts_str = format_timestamp_filename(slide.timestamp_sec)
            filename = f"{self.config.slide_prefix}_{slide.index:03d}_{ts_str}.png"
            filepath = out_dir / filename

            if self.config.save_png:
                # Save full-resolution BGR frame as PNG
                cv2.imwrite(str(filepath), slide.frame_bgr)
                image_paths.append(filepath)

            metadata_records.append({
                "index": slide.index,
                "timestamp_seconds": round(slide.timestamp_sec, 2),
                "timestamp_formatted": slide.formatted_timestamp,
                "filename": filename,
                "ssim_score": round(slide.ssim_score, 4),
                "phash_distance": slide.phash_distance,
            })

        # 2. Save JSON Metadata
        metadata_path: Optional[Path] = None
        if self.config.save_metadata:
            metadata_path = out_dir / "slides_metadata.json"
            meta_data = {
                "total_slides": len(slides),
                "video_info": {
                    "filename": probe_info.file_path.name if probe_info else "unknown",
                    "duration_seconds": probe_info.duration_seconds if probe_info else None,
                    "duration_formatted": probe_info.formatted_duration if probe_info else None,
                    "width": probe_info.width if probe_info else (slides[0].frame_bgr.shape[1] if slides else 0),
                    "height": probe_info.height if probe_info else (slides[0].frame_bgr.shape[0] if slides else 0),
                    "fps": probe_info.native_fps if probe_info else None,
                },
                "slides": metadata_records
            }
            with open(metadata_path, "w", encoding="utf-8") as f:
                json.dump(meta_data, f, indent=2)

        # 3. Export PDF
        pdf_path: Optional[Path] = None
        if self.config.save_pdf and slides:
            pdf_path = out_dir / "presentation.pdf"
            self._generate_pdf(slides, pdf_path)

        # 4. Export PPTX
        pptx_path: Optional[Path] = None
        if self.config.save_pptx and slides:
            pptx_path = out_dir / "presentation.pptx"
            self._generate_pptx(slides, image_paths, pptx_path, out_dir)

        # 5. Export ZIP Archive
        zip_path: Optional[Path] = None
        if self.config.save_zip and image_paths:
            zip_path = out_dir / "slides.zip"
            self._generate_zip(image_paths, metadata_path, zip_path)

        return ExportResult(
            output_dir=out_dir,
            total_slides=len(slides),
            image_paths=image_paths,
            metadata_path=metadata_path,
            pdf_path=pdf_path,
            pptx_path=pptx_path,
            zip_path=zip_path,
        )

    def _generate_pdf(self, slides: List[DetectedSlide], output_pdf: Path):
        """Combine slide frames into a single PDF document using Pillow."""
        pil_images = []
        for slide in slides:
            # Convert BGR OpenCV image to RGB PIL image
            rgb_frame = cv2.cvtColor(slide.frame_bgr, cv2.COLOR_BGR2RGB)
            pil_img = Image.fromarray(rgb_frame)
            pil_images.append(pil_img)

        if pil_images:
            first_img = pil_images[0]
            rest_imgs = pil_images[1:]
            first_img.save(
                output_pdf,
                "PDF",
                resolution=100.0,
                save_all=True,
                append_images=rest_imgs
            )

    def _generate_pptx(
        self,
        slides: List[DetectedSlide],
        image_paths: List[Path],
        output_pptx: Path,
        temp_dir: Path
    ):
        """Create a PowerPoint deck with full-bleed slide images."""
        prs = Presentation()

        # Determine aspect ratio from first slide
        h, w = slides[0].frame_bgr.shape[:2]
        aspect_ratio = w / max(1, h)

        if aspect_ratio >= 1.6:
            # 16:9 widescreen
            prs.slide_width = Inches(13.333)
            prs.slide_height = Inches(7.5)
        else:
            # 4:3 standard
            prs.slide_width = Inches(10)
            prs.slide_height = Inches(7.5)

        blank_slide_layout = prs.slide_layouts[6]  # Blank slide layout in standard template

        for i, slide in enumerate(slides):
            # Use saved image file if available, otherwise write temporary PNG
            if i < len(image_paths) and image_paths[i].exists():
                img_file = image_paths[i]
                cleanup = False
            else:
                img_file = temp_dir / f"_temp_pptx_{i}.png"
                cv2.imwrite(str(img_file), slide.frame_bgr)
                cleanup = True

            prs_slide = prs.slides.add_slide(blank_slide_layout)
            prs_slide.shapes.add_picture(
                str(img_file),
                left=0,
                top=0,
                width=prs.slide_width,
                height=prs.slide_height
            )

            if cleanup and img_file.exists():
                img_file.unlink()

        prs.save(str(output_pptx))

    def _generate_zip(
        self,
        image_paths: List[Path],
        metadata_path: Optional[Path],
        output_zip: Path
    ):
        """Package all images and metadata into a ZIP file."""
        with zipfile.ZipFile(output_zip, "w", zipfile.ZIP_DEFLATED) as zf:
            for img_path in image_paths:
                if img_path.exists():
                    zf.write(img_path, arcname=img_path.name)
            if metadata_path and metadata_path.exists():
                zf.write(metadata_path, arcname=metadata_path.name)
